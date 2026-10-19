# -*- coding: utf-8 -*-
import os, sys, collections, datetime
sys.path.insert(0, os.path.dirname(__file__))
from content import *
from docx_lib import Informe, construir
from audit import ejecutar

BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
out = lambda *p: os.path.join(BASE, *p)
RES, MET = ejecutar()
N_OK = sum(1 for r in RES if r[3]); N_TOT = len(RES)

sev = collections.Counter((c[1], c[3]) for c in CAMBIOS)
tot = collections.Counter(c[1] for c in CAMBIOS)
prom = {k: round(sum(v) / len(v), 1) for k, v in PUNTAJES.items()}
n_crit = sum(1 for c in CAMBIOS if c[3] == "Crítico")
n_may = sum(1 for c in CAMBIOS if c[3] == "Mayor")
n_men = sum(1 for c in CAMBIOS if c[3] == "Menor")


import re
def etiqueta(ref):
    m = re.match(r"^(.*?)\s*\((?:[^()]*?)(\d{4}|s\. f\.[^)]*)", ref)
    if not m: return ref[:40]
    aut = m.group(1).split(",")[0].replace("Ley 19.628 de 1999. Sobre protección de la vida privada. Diario Oficial de la República de Chile.", "Ley 19.628")
    return f"{aut} ({m.group(2)})"


def portada(extra_lineas):
    return dict(autor=AUTOR, lineas=[("Autor", AUTOR), ("Unidad", UNIDAD),
                ("Proyecto", "Ingeniería de prompts para diseño instruccional y auditoría de IA"), ("Fecha", FECHA)] + extra_lineas)


def cierre_ia(inf):
    inf.nota(f"Este documento se elaboró con asistencia de {HERRAMIENTA}, el {FECHA}, para analizar, redactar y sistematizar. "
             "La verificación humana de contenido y fuentes está pendiente de realización por el autor, quien asume la responsabilidad final. "
             "Detalle completo en el archivo 05_Declaracion_Uso_IA.", "Declaración de uso de IA")


def tabla_antes_despues(inf):
    rows = []
    for i, c in enumerate(CRITERIOS_Q):
        rows.append([c] + [PUNTAJES[k][i] for k in ("P1 original", "P1 v3.0", "P2 original", "P2 v3.0")])
    rows.append(["**Promedio**"] + [f"**{prom[k]}**".replace(".", ",") for k in ("P1 original", "P1 v3.0", "P2 original", "P2 v3.0")])
    inf.tabla(["Criterio (escala 1–5)", "P1 original", "P1 v3.0", "P2 original", "P2 v3.0"], rows, [2.8, 0.9, 0.9, 0.95, 0.95], 9.5,
              "Fuente: elaboración propia (autoevaluación asistida por IA; no es una medición empírica).")


def refs_lista(inf, cats=None):
    orden = ["Académica", "Norma de citación", "Normativa", "Institucional", "Material analizado", "Técnica / IA"]
    titulos = {"Académica": "Fuentes académicas", "Norma de citación": "Norma de citación", "Normativa": "Fuentes normativas",
               "Institucional": "Fuentes institucionales", "Material analizado": "Material analizado", "Técnica / IA": "Fuentes técnicas y herramientas de IA"}
    for cat in orden:
        items = [r for r in REFS if r[1] == cat and (cats is None or cat in cats)]
        if not items: continue
        inf.h2(titulos[cat])
        for r in items:
            par = inf.doc.add_paragraph(); par.paragraph_format.left_indent = Inches_(0.5); par.paragraph_format.first_line_indent = Inches_(-0.5)
            par.paragraph_format.line_spacing = 1.5; par.paragraph_format.space_after = Pt_(6)
            from docx_lib import rich
            rich(par, r[0], 11)


from docx.shared import Inches as Inches_, Pt as Pt_


# =============== 01 INFORME EJECUTIVO ===============
def f_ejecutivo():
    inf = Informe("Informe ejecutivo de optimización y auditoría de prompts", portada([("Producto", "01 · Informe ejecutivo")]),
                  "Informe ejecutivo · Optimización de prompts")
    inf.h1("1. Introducción")
    inf.p("Este informe resume el análisis, la optimización y la auditoría de dos instrucciones de sistema (prompts) del autor: el «Arquitecto Pedagógico de Casos Situados» (P1, versión 2.0), orientado a co-diseñar casos de estudio para educación técnico-profesional, y el «Prompt maestro de auditoría científica de resultados de IA» (P2), orientado a contrastar dos informes generados por IA sobre un mismo conjunto de datos.")
    inf.p("El objetivo fue elevar ambos prompts a una versión 3.0 más precisa, verificable y reutilizable, alineada con los estándares de presentación de Duoc UC, con citación APA 7.ª edición y con declaración explícita del uso de IA. El informe se organiza en hallazgos, mejoras, resultados, riesgos y recomendaciones.")
    inf.h1("2. Desarrollo")
    inf.h2("2.1 Hallazgos principales")
    inf.p(f"Se identificaron **{len(CAMBIOS)} oportunidades de mejora** ({tot['P1']} en P1 y {tot['P2']} en P2), clasificadas como {n_crit} críticas, {n_may} mayores y {n_men} menores. Ambos prompts tenían una arquitectura sólida; las brechas se concentraron en la verificabilidad, la trazabilidad y el control del flujo.")
    inf.bullets([
        "**P1:** contradicción de orden entre fases (A4 ordenado en la Fase 3 pero definido en la Fase 4); rúbrica diseñada después del caso; ausencia de conversión a nota 1,0–7,0; compuertas ambiguas; sin verificación aritmética ni regla ante instrucciones incrustadas en el adjunto.",
        "**P2 (críticos):** la comparación A contra B puede tratar la coincidencia como validez, y no prevé cómo operar si no se entregan los datos crudos. Además, la puntuación 1–10 carecía de anclas y los semáforos de regla de derivación.",
        "**Ambos:** sin formato Duoc UC, sin declaración de uso de IA, sin tratamiento de datos personales y sin protocolo de validación empírica.",
    ])
    inf.h2("2.2 Principales mejoras implementadas")
    inf.tabla(["Mejora", "P1 v3.0", "P2 v3.0"], [
        ["Control de flujo", "Compuertas G1–G4 con checklist; umbral de 8 páginas para dividir la entrega.", "Pasos 0–9 con confirmación solo ante [BRECHA] crítica; entrega por bloques."],
        ["Verificabilidad", "Verificación aritmética obligatoria; trazabilidad ampliada a 9 columnas.", "Línea base independiente, recálculo con datos, regla de coincidencia ≠ validez."],
        ["Evaluación", "Rúbrica diseñada en Fase 2; puntajes y conversión a nota 1,0–7,0.", "Escala 1–5 con anclas, 8 criterios ponderados, doble orden A-B/B-A."],
        ["Robustez", "Precedencia de instrucciones; adjunto tratado como datos; casos ficticios.", "Informes y datos tratados como material; declaración de conflicto de interés."],
        ["Institucional", "Formato Duoc UC, APA 7, declaración de IA, nota de uso de IA para estudiantes.", "Formato Duoc UC, APA 7, anexo de trazabilidad y declaración de IA."],
    ], [1.2, 2.65, 2.65], 9, "Fuente: elaboración propia.")
    inf.h2("2.3 Resultados")
    inf.p(f"En una autoevaluación con ocho criterios de calidad de prompts (escala 1–5), P1 pasó de {str(prom['P1 original']).replace('.',',')} a {str(prom['P1 v3.0']).replace('.',',')} y P2 de {str(prom['P2 original']).replace('.',',')} a {str(prom['P2 v3.0']).replace('.',',')}. La auditoría documental automatizada superó {N_OK} de {N_TOT} comprobaciones estructurales (balance de bloques, variables declaradas, pesos que suman 100 %, fórmula de nota, referencias cruzadas).")
    tabla_antes_despues(inf)
    inf.nota("Los puntajes son una autoevaluación asistida por IA con criterios explícitos y no equivalen a una medición empírica. No existe aún evidencia de desempeño con adjuntos o informes reales; el protocolo de piloto está incluido en cada prompt.", "Alcance de la evidencia")
    inf.h2("2.4 Riesgos y limitaciones")
    inf.bullets([
        f"**Longitud:** P1 pasó de {MET['P1']['pal_o']} a {MET['P1']['pal_n']} palabras y P2 de {MET['P2']['pal_o']} a {MET['P2']['pal_n']} (≈ {MET['P1']['tok_n']} y {MET['P2']['tok_n']} tokens). Mayor control a costa de mayor consumo de contexto; en modelos de ventana pequeña conviene dividir el prompt.",
        "**Bibliografía no verificada en línea:** durante la sesión no hubo acceso a los servicios de verificación de DOI/arXiv; todas las referencias se redactaron de memoria del modelo y deben comprobarse antes de uso formal (prioridad alta: Zheng, M. et al., 2024).",
        "**Normativa y formato institucional:** el título exacto y la URL de la norma de Bibliotecas Duoc UC, así como la vigencia de la Ley 21.719, deben confirmarse. El logotipo oficial debe insertarse manualmente en las portadas.",
        "**Convención de nota:** la fórmula lineal por tramos con 4,0 al 70 % debe validarse con la normativa de evaluación de la sede.",
    ])
    inf.h2("2.5 Recomendaciones y próximos pasos")
    inf.tabla(["N.º", "Acción", "Responsable sugerido", "Plazo sugerido"], [
        ["1", "Verificar cada referencia (DOI, arXiv, títulos) y completar URL normativas.", "Autor", "Antes de cualquier difusión"],
        ["2", "Ejecutar el piloto de P1 con tres adjuntos de distinta dificultad y dos revisores pares.", "Autor y pares", "2 semanas"],
        ["3", "Ejecutar el piloto de P2 con tres casos de errores sembrados y calcular sensibilidad y kappa.", "Autor y analista de datos", "2 semanas"],
        ["4", "Confirmar la conversión a nota y la plantilla de portada con la Escuela.", "Autor y jefatura", "1 semana"],
        ["5", "Versionar los prompts (v3.1) con los resultados del piloto.", "Autor", "Tras el piloto"],
    ], [0.45, 3.45, 1.45, 1.15], 9, "Fuente: elaboración propia. Responsables y plazos son propuestas.")
    inf.h1("3. Conclusión")
    inf.p("Las versiones 3.0 corrigen las brechas detectadas en cobertura, coherencia y verificabilidad, y cumplen el objetivo de entregar prompts documentados, trazables y alineados con el formato y las exigencias de declaración de IA de Duoc UC. La mejora es verificable en lo estructural y documental; su efecto sobre la calidad de los productos finales es una hipótesis que el protocolo de piloto permite contrastar. Hasta entonces, las referencias y la convención de nota requieren confirmación humana.")
    inf.h1("Referencias")
    refs_lista(inf, ["Académica", "Norma de citación", "Normativa", "Institucional", "Material analizado", "Técnica / IA"])
    inf.h1("Anexo A. Declaración de uso de IA")
    cierre_ia(inf)
    return inf


# =============== 02 INFORME TÉCNICO ===============
def f_tecnico():
    inf = Informe("Informe técnico de análisis y auditoría de prompts", portada([("Producto", "02 · Informe técnico")]),
                  "Informe técnico · Optimización de prompts")
    inf.h1("1. Introducción")
    inf.p("El informe documenta el procedimiento, los criterios y los resultados del análisis de dos prompts de uso académico. Su finalidad es que cada decisión de cambio pueda ser auditada: qué se observó, por qué constituye una brecha, qué se modificó y con qué fundamento.")
    inf.p("**Objetivos.** (1) Diagnosticar P1 y P2 según su naturaleza y propósito. (2) Optimizarlos en versiones 3.0 con criterios de ingeniería de prompts, diseño instruccional y práctica de auditoría. (3) Sistematizar la evidencia y la trazabilidad de cambios. (4) Auditar los resultados y declarar el uso de IA y las fuentes.")
    inf.p("**Organización.** El capítulo 2 presenta el marco metodológico; los capítulos 3 y 4, los diagnósticos; el 5, el diseño de las versiones 3.0; el 6, los resultados; el 7, la auditoría de calidad; el 8, las limitaciones; y el 9, las conclusiones.")
    inf.h1("2. Marco metodológico")
    inf.h2("2.1 Materiales analizados")
    inf.tabla(["Código", "Documento", "Naturaleza y propósito", "Tamaño"], [
        ["P1", "Instrucción de sistema «GEM — Arquitecto Pedagógico de Casos Situados», v2.0", "Prompt de co-diseño instruccional con interacción por fases y aprobación docente.", f"{MET['P1']['pal_o']} palabras"],
        ["P2", "Prompt maestro de auditoría científica de resultados de IA", "Prompt de evaluación y síntesis ejecutiva de dos informes de IA sobre los mismos datos.", f"{MET['P2']['pal_o']} palabras"],
    ], [0.6, 2.2, 2.9, 0.8], 9, "Fuente: archivos adjuntos por el autor.")
    inf.h2("2.2 Procedimiento")
    inf.bullets([
        "**Lectura crítica estructurada** de cada prompt: rol, entradas, procedimiento, criterios de calidad, salidas y manejo de excepciones.",
        "**Detección de brechas** con lista de verificación de ingeniería de prompts: ambigüedad, redundancia, contradicción, omisión, sesgo inducido, falta de control de flujo, falta de operacionalización y de trazabilidad.",
        "**Clasificación por severidad:** *Crítico* (puede producir resultados erróneos o dañinos sin detección), *Mayor* (degrada calidad, seguridad o auditabilidad) y *Menor* (mejora de precisión o usabilidad).",
        "**Rediseño** conservando lo que funcionaba y registrando cada cambio con identificador, fundamento y ubicación (Matriz de trazabilidad, producto 03).",
        "**Auditoría documental automatizada** de las versiones 3.0 (comprobaciones objetivas) y **autoevaluación** con rúbrica de ocho criterios.",
    ])
    inf.h2("2.3 Instrumento de evaluación de calidad de prompts")
    inf.p("Ocho criterios, escala 1–5 con anclas: 1 = ausente o deficiente; 2 = débil; 3 = aceptable; 4 = sólido; 5 = ejemplar y verificable. Los criterios se derivan de las buenas prácticas de ingeniería de prompts (rol y objetivo, entradas, criterios explícitos, validación), de seguridad (anti-alucinación, instrucciones incrustadas) y de las exigencias institucionales.")
    inf.tabla(["N.º", "Criterio", "Qué se observa"], [[str(i + 1), c, d] for i, (c, d) in enumerate(zip(CRITERIOS_Q, [
        "Rol, objetivo y destinatario inequívocos.", "Variables requeridas, valores por defecto y qué hacer si faltan.",
        "Estándares de calidad medibles antes de entregar.", "Compuertas, autocontrol, protocolos de piloto.",
        "Defensas ante alucinación, complacencia e instrucciones incrustadas; manejo de datos personales.",
        "Formato Duoc UC, APA 7, declaración de IA.", "Anclas, reglas deterministas, formatos de salida concretos.",
        "Versionado, documentación, trazabilidad y portabilidad."]))], [0.4, 2.6, 3.5], 9.5, "Fuente: elaboración propia.")
    inf.h1("3. Diagnóstico del Prompt 1 (Arquitecto Pedagógico de Casos)")
    inf.h2("3.1 Fortalezas conservadas")
    inf.bullets(["Rol acotado y propósito claro (comprender y usar el contenido, no memorizar).",
                 "Fidelidad al adjunto con marcador [SUGERENCIA EXTERNA] y [BRECHA].",
                 "Estructura narrativa de tres actos y cuatro desafíos graduados (D1–D4).",
                 "Cajas de apoyo, rúbrica de cuatro niveles y suite docente con cronograma.",
                 "Manejo de situaciones especiales, incluida la advertencia en contenidos de alto riesgo."])
    inf.h2("3.2 Brechas detectadas")
    filas = [[c[0], c[3], c[4], c[5]] for c in CAMBIOS if c[1] == "P1"]
    inf.tabla(["ID", "Sev.", "Hallazgo en el original", "Cambio en v3.0"], filas, [0.55, 0.6, 2.7, 2.65], 8.5, "Fuente: elaboración propia (detalle y fundamentos en el producto 03).")
    inf.h1("4. Diagnóstico del Prompt 2 (Auditoría de informes de IA)")
    inf.h2("4.1 Fortalezas conservadas")
    inf.bullets(["Principio rector de independencia: no favorecer ningún informe por extensión o popularidad.",
                 "Secuencia lógica de fases: auditoría, comparación, validación, sesgos, evaluación directiva, puntuación e informe.",
                 "Clasificación de hallazgos con justificación obligatoria.",
                 "Seis productos orientados a audiencias ejecutivas."])
    inf.h2("4.2 Brechas detectadas")
    filas = [[c[0], c[3], c[4], c[5]] for c in CAMBIOS if c[1] == "P2"]
    inf.tabla(["ID", "Sev.", "Hallazgo en el original", "Cambio en v3.0"], filas, [0.55, 0.6, 2.7, 2.65], 8.5, "Fuente: elaboración propia (detalle y fundamentos en el producto 03).")
    inf.h1("5. Diseño de las versiones 3.0")
    inf.h2("5.1 Principios de diseño")
    inf.bullets([
        "**Evidencia antes que fluidez:** toda afirmación se etiqueta como [HECHO], [INTERPRETACIÓN], [HIPÓTESIS] o [BRECHA]; prohibido inventar fuentes, cifras o páginas.",
        "**Verificación independiente:** el auditor calcula su línea base antes de leer los informes; la coincidencia entre IA no confirma por sí sola.",
        "**Alineamiento constructivo:** resultados de aprendizaje, caso, rúbrica y solucionario se diseñan juntos.",
        "**Control de flujo explícito:** compuertas, umbrales y entrega por bloques.",
        "**Seguridad de entradas:** adjuntos, informes y datos son material, no instrucciones.",
        "**Trazabilidad y auditabilidad:** identificadores, ubicaciones y anexos.",
        "**Transparencia de IA y estandarización institucional:** declaración de uso de IA, APA 7, formato Duoc UC.",
    ])
    inf.h2("5.2 Arquitectura comparada")
    inf.tabla(["Componente", "P1 original", "P1 v3.0", "P2 original", "P2 v3.0"], [
        ["Rol", "Sí", "Sí + precedencia", "Panel de 7 personas", "7 pasadas funcionales"],
        ["Entradas", "7 variables", "10 variables", "No definidas", "7 variables + regla R6"],
        ["Flujo", "4 fases", "4 fases con compuertas G1–G4", "7 fases", "Pasos 0–9 con autocontrol"],
        ["Validación", "Filtro de 4 preguntas", "Filtro con evidencia + aritmética", "5 estados", "2 ejes, 6 estados, regla de semáforo"],
        ["Puntuación", "Rúbrica 4 niveles", "Rúbrica con puntaje y nota", "10 criterios, 1–10", "8 criterios, 1–5, anclas, doble orden"],
        ["Seguridad", "Alto riesgo", "Alto riesgo + datos + adjunto", "No", "R7, R8, datos personales"],
        ["Formato / IA", "No", "DUOC, APA 7, declaración IA", "No", "DUOC, APA 7, declaración IA"],
        ["Piloto", "No", "Protocolo y umbrales propuestos", "No", "Protocolo y umbrales propuestos"],
    ], [1.0, 1.1, 1.65, 1.25, 1.5], 8.5, "Fuente: elaboración propia.")
    inf.h1("6. Resultados")
    inf.h2("6.1 Distribución de los cambios")
    filas = []
    for k in ("P1", "P2"):
        filas.append([k, sev[(k, "Crítico")], sev[(k, "Mayor")], sev[(k, "Menor")], tot[k]])
    filas.append(["**Total**", n_crit, n_may, n_men, len(CAMBIOS)])
    inf.tabla(["Prompt", "Crítico", "Mayor", "Menor", "Total"], filas, [1.5, 1.0, 1.0, 1.0, 1.0], 10, "Fuente: matriz de trazabilidad de cambios.")
    inf.h2("6.2 Autoevaluación de calidad")
    tabla_antes_despues(inf)
    inf.p("**Interpretación.** [INTERPRETACIÓN] El mayor avance se observa en P2, cuyo diseño original carecía de definición de entradas, de anclas de puntuación y de reglas de derivación; en P1 el avance proviene sobre todo del alineamiento constructivo, la conversión de nota y la seguridad de entradas. Al ser una autoevaluación realizada con asistencia de IA y por el mismo agente que rediseñó los prompts, existe riesgo de sobreestimación; por ello se incluye el protocolo de piloto.")
    inf.h1("7. Auditoría de control de calidad")
    inf.h2("7.1 Comprobaciones documentales automatizadas")
    filas = [[r[0], r[1], r[2], "Cumple" if r[3] else "NO CUMPLE", r[4]] for r in RES]
    inf.tabla(["ID", "Prompt", "Comprobación", "Resultado", "Evidencia"], filas, [0.55, 0.5, 2.6, 0.75, 2.1], 8, f"Fuente: script de auditoría reproducible (_build/audit.py). Resultado: {N_OK} de {N_TOT} comprobaciones cumplidas.")
    inf.nota("Las comprobaciones son estructurales y de coherencia interna. No demuestran que el prompt produzca mejores resultados con un modelo real; eso requiere el piloto descrito en cada prompt.", "Alcance")
    inf.h2("7.2 Filtro de calidad (FY)")
    inf.tabla(["Dimensión", "Pregunta", "Resultado", "Evidencia"], [
        ["Cobertura", "¿Se abordaron análisis, optimización, sistematización, bibliografía/IA, estandarización, auditoría, prompts y entregables?", "Sí", "Productos 01–08 y 41 cambios trazados."],
        ["Coherencia", "¿Objetivos, cambios, resultados y conclusiones se corresponden?", "Sí", "Cada cambio tiene ID, fundamento y sección de v3.0."],
        ["Aplicabilidad", "¿Es viable con recursos del docente y sin herramientas de compilación?", "Sí, con reserva", "Archivos TXT/DOCX/XLSX; la longitud de los prompts exige modelos de contexto amplio."],
        ["Accesibilidad", "¿Es comprensible para el destinatario?", "Sí", "Lenguaje formal, tablas y resumen ejecutivo breve."],
    ], [1.0, 2.6, 0.9, 2.0], 9, "Fuente: elaboración propia.")
    inf.h1("8. Limitaciones")
    inf.bullets([
        "**[BRECHA]** Sin ensayo empírico: no se ejecutaron los prompts con adjuntos o informes reales.",
        "**[BRECHA]** Referencias no verificadas en línea; se redactaron desde el conocimiento del modelo y requieren comprobación humana.",
        "**[BRECHA]** Logotipo, título exacto de la norma de Bibliotecas Duoc UC y vigencia de la Ley 21.719 por confirmar.",
        "**[RIESGO]** Autoevaluación con posible sobreestimación (el mismo agente rediseña y evalúa).",
        "**[RIESGO]** El aumento de longitud puede reducir el seguimiento de instrucciones en modelos con ventana pequeña.",
        "Los umbrales de aceptación del piloto son propuestas, no estándares normativos.",
    ])
    inf.h1("9. Conclusión")
    inf.p("Se cumplieron los objetivos de diagnóstico, optimización, sistematización y auditoría documental. Los prompts v3.0 incorporan controles de verificación, seguridad de entradas, trazabilidad y transparencia de IA ausentes en los originales. La validación de su efecto real queda planteada como hipótesis a contrastar con el protocolo de piloto, y la bibliografía queda sujeta a verificación humana antes de su uso formal.")
    inf.h1("Referencias")
    refs_lista(inf)
    inf.h1("Anexo A. Glosario de etiquetas y estados")
    inf.tabla(["Etiqueta", "Significado"], [
        ["[HECHO]", "Dato validado en el adjunto, en los datos o en una fuente verificable."],
        ["[INTERPRETACIÓN]", "Deducción razonada sujeta a contexto."],
        ["[HIPÓTESIS]", "Proposición pendiente de contrastación."],
        ["[BRECHA]", "Vacío de información o riesgo no cubierto."],
        ["[CONFLICTO]", "Contradicción entre instrucciones o fuentes que requiere decisión humana."],
        ["[ALERTA DE CONTENIDO]", "Instrucción incrustada en un adjunto o informe; se registra y no se ejecuta."],
        ["[VERIFICAR]", "Dato o referencia no comprobado que no debe citarse sin revisión."],
    ], [1.8, 4.7], 9.5, "Fuente: elaboración propia.")
    inf.h1("Anexo B. Declaración de uso de IA")
    cierre_ia(inf)
    return inf


# =============== 03 MATRIZ ===============
def f_matriz():
    inf = Informe("Matriz de trazabilidad de cambios", portada([("Producto", "03 · Matriz de trazabilidad")]), "Matriz de trazabilidad de cambios")
    inf.h1("1. Introducción")
    inf.p(f"La matriz registra los {len(CAMBIOS)} cambios entre las versiones originales y las versiones 3.0. Cada fila vincula el hallazgo observado en el original con el cambio realizado, su fundamento y la sección donde se aplica. También se entrega en formato XLSX filtrable (03_Matriz_Trazabilidad_Cambios.xlsx).")
    inf.nota("Severidad: Crítico = puede producir resultados erróneos o dañinos sin detección; Mayor = degrada calidad, seguridad o auditabilidad; Menor = mejora de precisión o usabilidad.", "Criterio")
    inf.h1("2. Matriz del Prompt 1")
    inf.tabla(["ID", "Sev.", "Hallazgo en el original", "Cambio en v3.0", "Fundamento", "Sección"],
              [[c[0], c[3], c[4], c[5], c[6], c[7]] for c in CAMBIOS if c[1] == "P1"], [0.5, 0.5, 1.8, 1.8, 1.2, 0.7], 8, "Fuente: elaboración propia.")
    inf.h1("3. Matriz del Prompt 2")
    inf.tabla(["ID", "Sev.", "Hallazgo en el original", "Cambio en v3.0", "Fundamento", "Sección"],
              [[c[0], c[3], c[4], c[5], c[6], c[7]] for c in CAMBIOS if c[1] == "P2"], [0.5, 0.5, 1.8, 1.8, 1.2, 0.7], 8, "Fuente: elaboración propia.")
    inf.h1("4. Elementos conservados sin cambios")
    inf.bullets(["P1: rol acotado, fidelidad al adjunto, tres actos, cuatro desafíos, cajas de apoyo, rúbrica de cuatro niveles, cronograma docente.",
                 "P2: principio de independencia, secuencia de fases, justificación obligatoria de clasificaciones, los seis productos ejecutivos."])
    inf.h1("Referencias")
    refs_lista(inf, ["Académica", "Norma de citación", "Normativa", "Institucional", "Material analizado", "Técnica / IA"])
    inf.h1("Anexo A. Declaración de uso de IA")
    cierre_ia(inf)
    return inf


# =============== 04 BIBLIOGRAFÍA ===============
def f_biblio():
    inf = Informe("Bibliografía normalizada y documentada", portada([("Producto", "04 · Bibliografía APA 7.ª edición")]), "Bibliografía normalizada")
    inf.h1("1. Introducción")
    inf.p("Este documento reúne las fuentes que sustentan el análisis y los prompts v3.0, normalizadas en APA 7.ª edición y diferenciadas en académicas, normativas, institucionales, material analizado y técnicas o de IA. Se indica además el uso de cada fuente y su estado de verificación.")
    inf.nota("Durante la elaboración no hubo acceso a servicios de verificación de DOI ni de arXiv. Todas las referencias deben comprobarse antes de su uso formal. Las que lo requieren están señaladas en el cuadro de verificación (capítulo 3). No se incorporaron fuentes cuya autoría o año no pudieran atribuirse con razonable certeza.", "Advertencia de verificación")
    inf.h1("2. Referencias")
    refs_lista(inf)
    inf.h1("3. Cuadro de uso y verificación")
    inf.tabla(["Fuente (autor, año)", "Categoría", "Uso en el proyecto", "Estado de verificación"],
              [[etiqueta(r[0]), r[1], r[2], r[3]] for r in REFS],
              [1.7, 1.0, 2.0, 1.8], 8.5, "Fuente: elaboración propia.")
    inf.h1("4. Criterios de citación de IA aplicados")
    inf.bullets([
        "La herramienta de IA se cita como software/modelo: Desarrollador. (año). *Nombre del modelo* (versión) [Modelo de lenguaje de gran tamaño]. URL.",
        "En el texto: (Anthropic, 2026).",
        "Las salidas de IA no se tratan como fuente de evidencia primaria; se declaran como asistencia en la Declaración de uso de IA.",
        "Los prompts y documentos del autor se citan como material inédito analizado.",
    ])
    inf.h1("Anexo A. Declaración de uso de IA")
    cierre_ia(inf)
    return inf


# =============== 05 DECLARACIÓN IA ===============
def f_ia():
    inf = Informe("Declaración de uso de inteligencia artificial", portada([("Producto", "05 · Declaración de uso de IA")]), "Declaración de uso de IA")
    inf.h1("1. Introducción")
    inf.p("Esta declaración transparenta el uso de herramientas de IA en el análisis, la optimización y la redacción de los productos del proyecto, conforme a las directrices académicas vigentes, e identifica la responsabilidad del autor.")
    inf.h1("2. Declaración")
    inf.tabla(["Elemento", "Detalle"], [
        ["Herramienta", f"{HERRAMIENTA}. Sesión de trabajo del {FECHA}."],
        ["Objetivo de utilización", "Analizar dos prompts, detectar brechas, proponer y redactar versiones 3.0, sistematizar resultados, elaborar informes en formato Duoc UC y organizar entregables."],
        ["Alcance de la asistencia", "La IA redactó los borradores de los prompts v3.0, los informes, la matriz, la bibliografía y los scripts de auditoría; ejecutó comprobaciones automatizadas. El autor aportó los prompts originales, las preferencias de formato y la dirección del trabajo."],
        ["Limitaciones declaradas", "Sin acceso a verificación en línea de DOI/arXiv; referencias redactadas desde el conocimiento del modelo. Sin ensayo empírico de los prompts. Autoevaluación realizada por el mismo agente que rediseñó."],
        ["Verificación humana", "PENDIENTE de realización por el autor: revisión de contenido, comprobación de referencias, confirmación de normativa, logotipo y convención de nota."],
        ["Responsabilidad final", "El autor asume la responsabilidad final por el contenido, su exactitud y su uso."],
        ["Datos personales", "No se procesaron datos personales de terceros."],
    ], [1.8, 4.7], 10, "Fuente: elaboración propia.")
    inf.h1("3. Texto sugerido para incluir en productos derivados")
    inf.nota(f"«En la elaboración de este documento se utilizó {HERRAMIENTA}, con fecha {FECHA}, para análisis, redacción y sistematización. El autor revisó, corrigió y validó el contenido y asume la responsabilidad final por su exactitud.»", "Plantilla")
    inf.p("El autor debe ajustar la frase al nivel real de revisión efectuado antes de firmar o difundir.")
    inf.h1("4. Cita de la herramienta")
    inf.p("Anthropic. (2026). *Claude Sonnet 5.5* (versión de octubre) [Modelo de lenguaje de gran tamaño]. https://claude.ai")
    inf.p("En el texto: (Anthropic, 2026).")
    inf.h1("Referencias")
    refs_lista(inf, ["Técnica / IA", "Norma de citación"])
    return inf


def txt_auditoria():
    L = []
    a = L.append
    a("=" * 80); a("AUDITORÍA DE LOS PROMPTS v3.0 — INFORME EN TEXTO PLANO"); a(f"Fecha: {FECHA} | Autor del proyecto: {AUTOR}")
    a(f"Asistencia de IA: {HERRAMIENTA}. Verificación humana: PENDIENTE."); a("=" * 80); a("")
    a("1. ALCANCE Y LÍMITES")
    a("- Auditoría documental: estructura, coherencia interna y cumplimiento de requisitos. Es reproducible (script _build/audit.py).")
    a("- NO es validación empírica: los prompts no se ejecutaron con adjuntos o informes reales.")
    a("- Referencias NO verificadas en línea (sin acceso a Crossref/arXiv). Verificar antes de uso formal.")
    a(""); a(f"2. COMPROBACIONES AUTOMATIZADAS ({N_OK}/{N_TOT} cumplidas)")
    for r in RES:
        a(f"[{'OK ' if r[3] else 'FALLA'}] {r[0]:6s} {r[1]:5s} {r[2]}"); a(f"        Evidencia: {r[4]}")
    a(""); a("3. MÉTRICAS DESCRIPTIVAS (original → v3.0)")
    for k, m in MET.items():
        a(f"- {k}: {m['pal_o']} → {m['pal_n']} palabras; {m['lin_o']} → {m['lin_n']} líneas; ≈{m['tok_n']} tokens (estimación: caracteres/4).")
    a(""); a("4. AUTOEVALUACIÓN DE CALIDAD (1–5) [INTERPRETACIÓN; no empírica]")
    a(f"{'Criterio':52s} P1-orig P1-v3 P2-orig P2-v3")
    for i, c in enumerate(CRITERIOS_Q):
        a(f"{c:52s} {PUNTAJES['P1 original'][i]:6d} {PUNTAJES['P1 v3.0'][i]:5d} {PUNTAJES['P2 original'][i]:6d} {PUNTAJES['P2 v3.0'][i]:5d}")
    a(f"{'PROMEDIO':52s} {prom['P1 original']:6} {prom['P1 v3.0']:5} {prom['P2 original']:6} {prom['P2 v3.0']:5}")
    a(""); a("5. CAMBIOS REGISTRADOS")
    a(f"- Total {len(CAMBIOS)}: {n_crit} críticos, {n_may} mayores, {n_men} menores (P1: {tot['P1']}; P2: {tot['P2']}). Detalle en 03_Matriz_Trazabilidad_Cambios.")
    a(""); a("6. FILTRO DE CALIDAD (FY)")
    a("- Cobertura: Sí. Coherencia: Sí. Aplicabilidad: Sí, con reserva (longitud). Accesibilidad: Sí.")
    a(""); a("7. RIESGOS RESIDUALES Y ACCIONES")
    for x in ["Longitud: dividir los prompts en módulos si el modelo pierde instrucciones (≈4.500–5.000 tokens cada uno).",
              "Verificar referencias (prioridad alta: Zheng, M. et al., 2024).",
              "Confirmar con la institución la fórmula de nota (4,0 al 70 %) y la norma de Bibliotecas Duoc UC.",
              "Ejecutar los protocolos de piloto incluidos en cada prompt y versionar a v3.1.",
              "Insertar el logotipo oficial en las portadas DOCX."]:
        a("- " + x)
    a(""); a("8. DECLARACIÓN DE USO DE IA")
    a(f"Herramienta: {HERRAMIENTA}. Objetivo: análisis, optimización y redacción. Alcance: borradores y comprobaciones; el autor dirige y valida. Responsabilidad final: del autor.")
    return "\n".join(L) + "\n"


def leeme():
    return f"""LÉEME — PAQUETE DE ENTREGABLES: PROMPTS v3.0 (optimizados y auditados)
Autor: {AUTOR} | {UNIDAD} | {FECHA}
Asistencia de IA: {HERRAMIENTA}. Verificación humana: PENDIENTE (ver 05_Declaracion_Uso_IA.docx).

CONTENIDO
01_Informe_Ejecutivo.docx ........ Hallazgos, mejoras, resultados, riesgos y próximos pasos (formato Duoc UC).
02_Informe_Tecnico.docx .......... Metodología, diagnóstico, diseño, resultados y auditoría completa.
03_Matriz_Trazabilidad_Cambios ... .docx y .xlsx con los {len(CAMBIOS)} cambios (ID, severidad, fundamento, sección).
04_Bibliografia_APA7.docx ........ Fuentes normalizadas, diferenciadas, con cuadro de verificación y uso de IA.
05_Declaracion_Uso_IA.docx ....... Herramienta, objetivo, alcance, verificación y responsabilidad.
06_Prompts_TXT/ .................. PROMPT_01 (Arquitecto Pedagógico de Casos) y PROMPT_02 (Auditor de informes de IA), v3.0.
07_Auditoria_Prompts/ ............ AUDITORIA_PROMPTS_v3.0.txt (comprobaciones, métricas, riesgos).
08_Originales/ ................... Prompts originales sin modificar.
_build/ .......................... Scripts reproducibles (contenido, auditoría y generación de documentos).

ANTES DE USAR O DIFUNDIR
1. Verificar cada referencia (DOI/arXiv) y completar URL normativas.
2. Insertar el logotipo oficial Duoc UC en las portadas.
3. Confirmar con la Escuela la conversión de nota y la norma de presentación.
4. Ejecutar los protocolos de piloto de cada prompt.
"""


def xlsx():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Cambios"
    head = ["ID", "Prompt", "Categoría", "Severidad", "Hallazgo en el original", "Cambio en v3.0", "Fundamento", "Sección v3.0"]
    ws.append(head)
    for c in CAMBIOS: ws.append(list(c))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Calibri"); c.fill = PatternFill("solid", fgColor="1A1A1A")
    for i, w in enumerate([8, 8, 24, 11, 60, 60, 40, 24], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].row % 2 == 1:
            for c in row: c.fill = PatternFill("solid", fgColor="EEEEEE")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    w2 = wb.create_sheet("Autoevaluación")
    w2.append(["Criterio"] + list(PUNTAJES.keys()))
    for i, c in enumerate(CRITERIOS_Q): w2.append([c] + [PUNTAJES[k][i] for k in PUNTAJES])
    n = len(CRITERIOS_Q) + 1
    w2.append(["Promedio"] + [f"=ROUND(AVERAGE({col}2:{col}{n}),1)" for col in "BCDE"])
    for c in w2[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1A1A1A")
    w2.column_dimensions["A"].width = 50
    w3 = wb.create_sheet("Auditoría")
    w3.append(["ID", "Prompt", "Comprobación", "Resultado", "Evidencia"])
    for r in RES: w3.append([r[0], r[1], r[2], "Cumple" if r[3] else "NO CUMPLE", r[4]])
    for c in w3[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1A1A1A")
    for col, w in zip("ABCDE", [8, 8, 70, 12, 60]): w3.column_dimensions[col].width = w
    wb.save(out("03_Matriz_Trazabilidad_Cambios.xlsx"))


if __name__ == "__main__":
    docs = [("01_Informe_Ejecutivo.docx", f_ejecutivo), ("02_Informe_Tecnico.docx", f_tecnico),
            ("03_Matriz_Trazabilidad_Cambios.docx", f_matriz), ("04_Bibliografia_APA7.docx", f_biblio),
            ("05_Declaracion_Uso_IA.docx", f_ia)]
    for nombre, fx in docs:
        paginas, n = construir(fx, out(nombre))
        print(nombre, "->", n, "páginas; índice:", len(paginas), "entradas")
    xlsx()
    open(out("07_Auditoria_Prompts", "AUDITORIA_PROMPTS_v3.0.txt"), "w", encoding="utf-8").write(txt_auditoria())
    open(out("00_LEEME.txt"), "w", encoding="utf-8").write(leeme())
