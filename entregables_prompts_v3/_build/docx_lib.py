# -*- coding: utf-8 -*-
"""Generador de DOCX con formato de informes DUOC UC (carta, Calibri, 1,5, paleta oficial)."""
import re, subprocess, os, shutil, tempfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NEGRO, AMARILLO, GRIS, GRIS_CLARO = "1A1A1A", "FFB800", "666666", "EEEEEE"
FONT = "Calibri"


def _rgb(h): return RGBColor.from_string(h)


def _shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = kw.get(edge)
        el = OxmlElement(f"w:{edge}")
        if spec:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(spec[0])); el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcPr.append(b)


def _par_border(par, edge, sz, color, space=4):
    pPr = par._p.get_or_add_pPr()
    pb = pPr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr"); pPr.append(pb)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
    pb.append(el)


def rich(par, text, size=None, color=None, bold=None):
    """Texto con *cursiva* y **negrita** simples."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok: continue
        if tok.startswith("**"): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*"): r = par.add_run(tok[1:-1]); r.italic = True
        else: r = par.add_run(tok)
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = _rgb(color)
        if bold is not None and not tok.startswith("**"): r.bold = bold if bold else r.bold
    return par


def _fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.5, before=0, after=6, keep_next=False):
    pf = par.paragraph_format
    pf.alignment = align; pf.line_spacing = line
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.keep_with_next = keep_next


def _page_field(par):
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = par.add_run()
        r.font.name = FONT; r.font.size = Pt(10); r.font.color.rgb = _rgb(GRIS)
        if t:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), t); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)


class Informe:
    def __init__(self, titulo, datos_portada, titulo_corto):
        self.doc = Document()
        self.titulo = titulo
        self.heads = []  # (nivel, texto)
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Inches(1))
        s.different_first_page_header_footer = True
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(titulo_corto + "  |  "); r.font.size = Pt(10); r.font.name = FONT; r.font.color.rgb = _rgb(GRIS)
        _page_field(fp)
        st = self.doc.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        cp = self.doc.core_properties
        cp.title = titulo.title(); cp.author = datos_portada["autor"]; cp.language = "es-CL"
        self._portada(datos_portada)

    # ---------- portada ----------
    def _portada(self, d):
        p = self.doc.add_paragraph(); _fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0, 0)
        rich(p, "[Insertar logotipo oficial Duoc UC en esta esquina]", size=9, color=GRIS)
        for _ in range(5):
            self.doc.add_paragraph()
        t = self.doc.add_paragraph(); _fmt(t, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 0, 10)
        r = t.add_run(self.titulo.upper()); r.bold = True; r.font.size = Pt(14); r.font.name = FONT; r.font.color.rgb = _rgb(NEGRO)
        _par_border(t, "bottom", 18, AMARILLO, 8)
        for _ in range(7):
            self.doc.add_paragraph()
        for etiqueta, valor in d["lineas"]:
            q = self.doc.add_paragraph(); _fmt(q, WD_ALIGN_PARAGRAPH.CENTER, 1.15, 0, 4)
            a = q.add_run(etiqueta + ": "); a.bold = True; a.font.size = Pt(14); a.font.name = FONT
            b = q.add_run(valor); b.font.size = Pt(14); b.font.name = FONT
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.idx_par = self.doc.add_paragraph()  # marcador del índice

    # ---------- bloques ----------
    def h1(self, txt):
        p = self.doc.add_paragraph(); _fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 14, 8, True)
        r = p.add_run(txt); r.bold = True; r.font.size = Pt(15); r.font.name = FONT; r.font.color.rgb = _rgb(NEGRO)
        _par_border(p, "left", 36, AMARILLO, 8)
        p.paragraph_format.left_indent = Inches(0.12)
        self.heads.append((1, txt))

    def h2(self, txt):
        p = self.doc.add_paragraph(); _fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 10, 6, True)
        r = p.add_run(txt); r.bold = True; r.font.size = Pt(12.5); r.font.name = FONT; r.font.color.rgb = _rgb(GRIS)
        self.heads.append((2, txt))

    def p(self, txt):
        par = self.doc.add_paragraph(); _fmt(par); rich(par, txt, 11)

    def bullets(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet"); _fmt(par, after=2)
            rich(par, it, 11)
        sp = self.doc.add_paragraph(); _fmt(sp, line=1.5, after=0)

    def caption(self, txt):
        par = self.doc.add_paragraph(); _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 6, 3, True)
        rich(par, txt, 10, GRIS, True)
        for r in par.runs: r.bold = True

    def nota(self, txt, titulo=None):
        t = self.doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; _shade(c, GRIS_CLARO)
        _cell_borders(c, left=(32, AMARILLO))
        par = c.paragraphs[0]; _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 3, 3)
        if titulo:
            r = par.add_run(titulo + ". "); r.bold = True; r.font.size = Pt(10.5); r.font.name = FONT
        rich(par, txt, 10.5)
        sp = self.doc.add_paragraph(); _fmt(sp, line=1.5, after=0)

    def tabla(self, header, rows, widths, size=9, fuente=None):
        t = self.doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for i, w in enumerate(widths): t.columns[i].width = Inches(w)
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.width = Inches(widths[i]); _shade(c, NEGRO)
            par = c.paragraphs[0]; _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 2, 2)
            r = par.add_run(h); r.bold = True; r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = _rgb("FFFFFF")
        trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for k, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].width = Inches(widths[i])
                if k % 2 == 1: _shade(cells[i], GRIS_CLARO)
                par = cells[i].paragraphs[0]; _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 2, 2)
                rich(par, str(v), size)
            cant = t.rows[-1]._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); cant.append(cs)
        if fuente:
            self.caption(fuente)
        else:
            sp = self.doc.add_paragraph(); _fmt(sp, line=1.5, after=0)

    def salto(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- índice ----------
    def _llenar_indice(self, paginas):
        """Inserta el índice (con números de página) en el marcador."""
        cur = self.idx_par
        t = OxmlElement("w:p"); cur._p.addprevious(t)
        from docx.text.paragraph import Paragraph
        par = Paragraph(t, cur._parent); _fmt(par, WD_ALIGN_PARAGRAPH.LEFT, 1.15, 0, 12)
        r = par.add_run("Índice"); r.bold = True; r.font.size = Pt(15); r.font.name = FONT
        _par_border(par, "left", 36, AMARILLO, 8); par.paragraph_format.left_indent = Inches(0.12)
        for (nivel, txt) in self.heads:
            q = OxmlElement("w:p"); cur._p.addprevious(q); pq = Paragraph(q, cur._parent)
            _fmt(pq, WD_ALIGN_PARAGRAPH.LEFT, 1.3, 0, 2)
            pq.paragraph_format.left_indent = Inches(0.0 if nivel == 1 else 0.3)
            pq.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            rr = pq.add_run(txt + "\t" + str(paginas.get(txt, "")))
            rr.font.name = FONT; rr.font.size = Pt(11); rr.bold = (nivel == 1)
        br = OxmlElement("w:p"); cur._p.addprevious(br); pb = Paragraph(br, cur._parent)
        pb.add_run().add_break(WD_BREAK.PAGE)

    def guardar(self, ruta, paginas=None):
        self._llenar_indice(paginas or {})
        self.doc.save(ruta)


def _pdf_pages(docx_path, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    return pdf, txt.split("\f")


def construir(factory, ruta_docx):
    """Dos pasadas: compone, mide la paginación con LibreOffice y vuelve a componer con índice real."""
    tmp = tempfile.mkdtemp()
    inf = factory(); borrador = os.path.join(tmp, "b.docx"); inf.guardar(borrador, {})
    _, pages = _pdf_pages(borrador, tmp)
    found = {}
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    for nivel, txt in inf.heads:
        for n, pg in enumerate(pages, 1):
            if n < 3: continue  # portada e índice
            if norm(txt) in norm(pg):
                found[txt] = n; break
    inf2 = factory(); inf2.guardar(ruta_docx, found)
    pdf, pages2 = _pdf_pages(ruta_docx, tmp)
    shutil.rmtree(tmp, ignore_errors=True)
    return found, len([x for x in pages2 if x.strip()])
