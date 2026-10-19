"""Librería de documentos con formato Duoc UC (norma Bibliotecas Duoc UC): carta, Calibri, 1,5, justificado,
paleta negro/amarillo/gris, portada, índice con páginas, Introducción-Desarrollo-Conclusión-Referencias-Anexos."""
import json, re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

BASE = '/home/user/Perfil_investigador/La_Puerta_Informes/'
R = json.load(open(BASE + 'src/resultados.json'))
G = BASE + 'graficos/'
# Paleta oficial Duoc UC (más tintes de amarillo/gris). Color de escuela Administración (#AC4FC6) no se usa.
NAVY, AMB, TEAL, CORAL, GRAY, LG, SKY = '1A1A1A', 'FFB800', '1A1A1A', 'FFB800', '666666', 'EEEEEE', '666666'
DARK_FILLS = ('1A1A1A', '666666')
FONT = 'Calibri'; W_TXT = 16.5
ORD = {
 'tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription'],
 'tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
 'pPr': ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']}
def put(parent, el, kind):
    order = ORD[kind]; tag = el.tag.split('}')[1]
    for old in parent.findall(qn('w:' + tag)): parent.remove(old)
    idx = order.index(tag)
    for ch in parent:
        t = ch.tag.split('}')[1]
        if t in order and order.index(t) > idx: ch.addprevious(el); return el
    parent.append(el); return el

def n(x, d=2): return f'{x:.{d}f}'.replace('.', ',')
def pc(x, d=0): return n(x, d) + ' %'
def pv(p): return 'p < 0,001' if p < 0.001 else 'p = ' + n(p, 3 if p < 0.01 else 2)
def text_on(fill): return 'FFFFFF' if fill in DARK_FILLS else NAVY

def shade(cell, hexcolor):
    put(cell._tc.get_or_add_tcPr(), parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hexcolor}"/>'), 'tcPr')
def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    put(cell._tc.get_or_add_tcPr(), m, 'tcPr')
def cell_border(cell, **sides):  # lado=(sz, color)
    b = OxmlElement('w:tcBorders')
    for k in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{k}')
        if k in sides: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sides[k][0])); e.set(qn('w:color'), sides[k][1])
        else: e.set(qn('w:val'), 'nil')
        b.append(e)
    put(cell._tc.get_or_add_tcPr(), b, 'tcPr')
def borders(table, color='CCCCCC', sz=4, inside=True):
    b = OxmlElement('w:tblBorders')
    for k in ['top', 'left', 'bottom', 'right'] + (['insideH', 'insideV'] if inside else []):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color); b.append(e)
    put(table._tbl.tblPr, b, 'tblPr')
def no_borders(table):
    b = OxmlElement('w:tblBorders')
    for k in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{k}'); e.set(qn('w:val'), 'nil'); b.append(e)
    put(table._tbl.tblPr, b, 'tblPr')
def set_col_widths(table, widths):
    tot = sum(widths)
    if tot > W_TXT: widths = [w * W_TXT / tot for w in widths]
    table.autofit = False
    tbl = table._tbl; tblPr = tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); put(tblPr, lay, 'tblPr')
    tw = OxmlElement('w:tblW'); tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(sum(widths) * 567))); put(tblPr, tw, 'tblPr')
    grid = tbl.tblGrid
    for gc in list(grid): grid.remove(gc)
    for w in widths:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w * 567))); grid.append(gc)
    for row in table.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
def keep_row(row): row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
def repeat_header(row):
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); row._tr.get_or_add_trPr().append(e)
def cellp(p, align='l', after=0, line=1.0):
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align == 'l' else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = line
def pborder(p, side, color=AMB, sz=36, space=8):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); e = OxmlElement(f'w:{side}')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color); b.append(e); put(pPr, b, 'pPr')

def runs(p, text, size=None, color=None, bold=None, italic=None):
    """Soporta **negrita**."""
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part: continue
        b = part.startswith('**') and part.endswith('**')
        r = p.add_run(part[2:-2] if b else part); r.font.name = FONT; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        r.bold = True if (b or bold) else None
        if italic: r.italic = True
    return p

class Doc:
    def __init__(self, titulo, subtitulo, codigo, cover_img, header_text):
        self.d = Document(); self.titulo = titulo; d = self.d; self.pn = 0; self.cn = 0; self.heads = []; self.header_text = header_text
        s = d.sections[0]; s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        s.left_margin = s.right_margin = Cm(2.54); s.top_margin = s.bottom_margin = Cm(2.54); s.header_distance = s.footer_distance = Cm(1.25)
        st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11); st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        st.font.color.rgb = RGBColor.from_string(NAVY)
        st.paragraph_format.space_after = Pt(8); st.paragraph_format.line_spacing = 1.5; st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for nm, sz, col in (('Heading 1', 16, NAVY), ('Heading 2', 13, NAVY), ('Heading 3', 12, GRAY), ('Heading 4', 11, GRAY)):
            h = d.styles[nm]; h.font.name = FONT; h.font.size = Pt(sz); h.font.bold = True; h.font.italic = False; h.font.color.rgb = RGBColor.from_string(col)
            rf = h.element.get_or_add_rPr().rFonts
            for a in ('ascii', 'hAnsi', 'eastAsia', 'cs'): rf.set(qn('w:' + a), FONT)
            for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                if rf.get(qn('w:' + a)) is not None: del rf.attrib[qn('w:' + a)]
            pf = h.paragraph_format; pf.space_before = Pt(16 if sz >= 13 else 10); pf.space_after = Pt(12); pf.keep_with_next = True; pf.line_spacing = 1.15; pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._cover(titulo, subtitulo, codigo)
        self.marker = d.add_paragraph(); self.marker.paragraph_format.space_after = Pt(0)
        r = self.marker.add_run(' '); r.font.size = Pt(1)
        hp = s.header.paragraphs[0]; runs(hp, header_text, 9, GRAY); cellp(hp)
        fp = s.footer.paragraphs[0]; cellp(fp, 'c'); fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; runs(fp, 'Duoc UC · Proyecto La Puerta · Página ', 9, GRAY); self._field(fp, 'PAGE')
    def _cover(self, titulo, subtitulo, codigo):
        d = self.d; s = d.sections[0]; s.different_first_page_header_footer = True
        t = d.add_table(rows=3, cols=1); no_borders(t); set_col_widths(t, [W_TXT])
        for row, hcm in zip(t.rows, (3.2, 10.6, 8.6)):
            trPr = row._tr.get_or_add_trPr(); hh = OxmlElement('w:trHeight'); hh.set(qn('w:val'), str(int(hcm * 567))); hh.set(qn('w:hRule'), 'exact'); trPr.append(hh)
        # 1) logo (marcador de posición: insertar logo oficial)
        c = t.rows[0].cells[0]; inner = c.add_table(rows=1, cols=1); no_borders(inner); set_col_widths(inner, [5.2]); inner.alignment = WD_TABLE_ALIGNMENT.LEFT; ic = inner.rows[0].cells[0]
        cell_border(ic, top=(6, '999999'), left=(6, '999999'), bottom=(6, '999999'), right=(6, '999999')); cell_margins(ic, 160, 160, 120, 120)
        p = ic.paragraphs[0]; cellp(p, 'c'); runs(p, '[Insertar logo Duoc UC]', 9, GRAY)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        # 2) título
        c = t.rows[1].cells[0]; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10); p.paragraph_format.line_spacing = 1.15
        self.cover_title = ('PROYECTO LA PUERTA: ' + titulo.upper()).rstrip('.'); runs(p, self.cover_title, 14, NAVY, bold=True)
        pborder(p, 'bottom', AMB, 24, 10)
        q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_before = Pt(10); q.paragraph_format.line_spacing = 1.15; runs(q, subtitulo, 12, GRAY)
        # 3) datos
        c = t.rows[2].cells[0]; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        lines = ['Autor y docente responsable: Juan Antonio Carrasco Navarrete', 'Escuela de Administración y Negocios · Sede San Bernardo', f'Proyecto La Puerta · {codigo.split(" · ")[0].title()}', 'Octubre de 2026']
        for i, ln in enumerate(lines):
            p = c.paragraphs[0] if i == 0 else c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15; runs(p, ln, 14, NAVY)
        p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(8); runs(p, 'Elaborado con apoyo de herramientas de Inteligencia Artificial (véase Anexos)', 9, GRAY)
    def _field(self, p, code):
        r = p.add_run(); r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(GRAY)
        for t, txt in (('begin', None), (None, code), ('end', None)):
            if t: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t)
            else: e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt
            r._r.append(e)
    # ---- estructura ----
    def part(self, title):
        self.pn += 1; self.cn = 0; txt = f'{self.pn}. {title}'
        h = self.d.add_heading(txt, 1); h.paragraph_format.page_break_before = True; pborder(h, 'left'); self.heads.append((1, txt)); return h
    def h1(self, t, newpage=False):
        t = re.sub(r'^\d+\.\s*', '', t); self.cn += 1; txt = f'{self.pn}.{self.cn} {t}'
        h = self.d.add_heading(txt, 2); pborder(h, 'left', AMB, 28, 8); self.heads.append((2, txt)); return h
    def h2(self, t): return self.d.add_heading(re.sub(r'^\d+(\.\d+)+\s+', '', t), 3)
    def h3(self, t): return self.d.add_heading(t, 4)
    def pagebreak(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    def p(self, t, size=None, color=None, italic=False, align=None, after=None):
        p = self.d.add_paragraph(); runs(p, t, size, color, italic=italic)
        if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None: p.paragraph_format.space_after = Pt(after)
        return p
    def bullets(self, items, size=None):
        for it in items:
            p = self.d.add_paragraph(style='List Bullet'); runs(p, it, size); p.paragraph_format.space_after = Pt(3)
    def numbered(self, items):
        for i, it in enumerate(items, 1):
            p = self.d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6); p.paragraph_format.space_after = Pt(4)
            runs(p, f'**{i}.**  ' + it)
    def _spacer(self, pts=10):
        p = self.d.add_paragraph(); p.paragraph_format.space_after = Pt(pts); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(' '); r.font.size = Pt(4)
    def callout(self, titulo, texto, color=None, bg=None, icon=''):
        t = self.d.add_table(rows=1, cols=1); no_borders(t); c = t.rows[0].cells[0]; shade(c, LG)
        cell_border(c, left=(32, AMB)); cell_margins(c, 120, 120, 220, 160); keep_row(t.rows[0])
        p = c.paragraphs[0]; cellp(p, 'l', 3, 1.15); runs(p, titulo, 10.5, NAVY, bold=True)
        for line in (texto if isinstance(texto, list) else [texto]):
            q = c.add_paragraph(); cellp(q, 'l', 3, 1.15); runs(q, line, 10)
        set_col_widths(t, [W_TXT]); self._spacer(10)
    def kpis(self, items):
        k = len(items); t = self.d.add_table(rows=1, cols=k); no_borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (v, lab, col) in enumerate(items):
            c = t.rows[0].cells[i]; shade(c, LG); cell_border(c, top=(24, AMB)); cell_margins(c, 140, 140, 120, 120)
            p = c.paragraphs[0]; cellp(p, 'c'); runs(p, v, 22, NAVY, bold=True)
            q = c.add_paragraph(); cellp(q, 'c'); runs(q, lab, 9, GRAY)
        set_col_widths(t, [W_TXT / k] * k); self._spacer(10)
    def table(self, header, rows, widths, hcolor=None, size=9, zebra=True, colors=None, align=None):
        t = self.d.add_table(rows=1, cols=len(header)); borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; shade(c, NAVY); cell_margins(c, 70, 70, 90, 90); p = c.paragraphs[0]; cellp(p); runs(p, h, size, 'FFFFFF', bold=True)
        repeat_header(t.rows[0])
        for ri, row in enumerate(rows):
            r = t.add_row(); keep_row(r)
            for i, v in enumerate(row):
                c = r.cells[i]; cell_margins(c, 60, 60, 90, 90); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if zebra and ri % 2: shade(c, LG)
                p = c.paragraphs[0]; cellp(p); col = None
                fill = colors.get((ri, i)) if colors else None
                if fill:
                    fill = {'1F8A5B': '1A1A1A', '8CBF3F': '666666'}.get(fill, fill); shade(c, fill); col = text_on(fill)
                runs(p, str(v), size, col, bold=bool(col))
                if align and align[i] == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_col_widths(t, widths); self._spacer(12); return t
    def fig(self, name, caption, width=16.0):
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(G + name + '.png', width=Cm(min(width, 16.0)))
        q = self.d.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(q, caption, 9, GRAY, italic=True); q.paragraph_format.space_after = Pt(12); q.paragraph_format.line_spacing = 1.15
    def lee(self, texto): self.callout('Cómo leer este gráfico', texto)
    # ---- índice y guardado ----
    def _index(self, pages):
        els = []
        def add(txt, level, page):
            p = self.d.add_paragraph(); pf = p.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(2); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.left_indent = Cm(0 if level == 1 else 0.8); pf.tab_stops.add_tab_stop(Cm(W_TXT), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            runs(p, txt, 11, NAVY, bold=(level == 1)); runs(p, '\t' + str(pages.get(txt, '—')), 11, NAVY, bold=(level == 1)); els.append(p._p)
        h = self.d.add_heading('Índice', 1); h.paragraph_format.page_break_before = True; pborder(h, 'left'); els.append(h._p)
        for lvl, txt in self.heads: add(txt, lvl, None)
        for e in els: self.marker._p.addprevious(e)
        self.marker._p.getparent().remove(self.marker._p)
    def save(self, path):
        base = os.path.basename(path)
        pf = BASE + 'src/pages_' + base + '.json'; pages = json.load(open(pf)) if os.path.exists(pf) else {}
        self._index(pages); json.dump(self.heads, open(BASE + 'src/heads_' + base + '.json', 'w'), ensure_ascii=False)
        self.d.save(path)

def cover_art(name, accent=None): return G + name + '.png'

BIB = [
 'Bandura, A. (1997). *Self-efficacy: The exercise of control*. W. H. Freeman.',
 'Braun, V., & Clarke, V. (2006). Using thematic analysis in psychology. *Qualitative Research in Psychology, 3*(2), 77–101.',
 'Campbell, D. T., & Stanley, J. C. (1963). *Experimental and quasi-experimental designs for research*. Rand McNally.',
 'Cliff, N. (1993). Dominance statistics: Ordinal analyses to answer ordinal questions. *Psychological Bulletin, 114*(3), 494–509.',
 'Cohen, J. (1988). *Statistical power analysis for the behavioral sciences* (2.ª ed.). Lawrence Erlbaum.',
 'Cronbach, L. J. (1951). Coefficient alpha and the internal structure of tests. *Psychometrika, 16*(3), 297–334.',
 'Dweck, C. S. (2006). *Mindset: The new psychology of success*. Random House.',
 'Granovetter, M. S. (1973). The strength of weak ties. *American Journal of Sociology, 78*(6), 1360–1380.',
 'Holm, S. (1979). A simple sequentially rejective multiple test procedure. *Scandinavian Journal of Statistics, 6*(2), 65–70.',
 'Kolb, D. A. (1984). *Experiential learning: Experience as the source of learning and development*. Prentice-Hall.',
 'Kruskal, W. H., & Wallis, W. A. (1952). Use of ranks in one-criterion variance analysis. *Journal of the American Statistical Association, 47*(260), 583–621.',
 'Lent, R. W., Brown, S. D., & Hackett, G. (1994). Toward a unifying social cognitive theory of career and academic interest, choice, and performance. *Journal of Vocational Behavior, 45*(1), 79–122.',
 'Mann, H. B., & Whitney, D. R. (1947). On a test of whether one of two random variables is stochastically larger than the other. *Annals of Mathematical Statistics, 18*(1), 50–60.',
 'Norman, G. (2010). Likert scales, levels of measurement and the “laws” of statistics. *Advances in Health Sciences Education, 15*(5), 625–632.',
 'Podsakoff, P. M., MacKenzie, S. B., Lee, J.-Y., & Podsakoff, N. P. (2003). Common method biases in behavioral research. *Journal of Applied Psychology, 88*(5), 879–903.',
 'Sullivan, G. M., & Artino, A. R. (2013). Analyzing and interpreting data from Likert-type scales. *Journal of Graduate Medical Education, 5*(4), 541–542.',
 'Tomlinson, M. (2017). Forms of graduate capital and their relationship to graduate employability. *Education + Training, 59*(4), 338–352.',
 'Vargha, A., & Delaney, H. D. (2000). A critique and improvement of the CL common language effect size statistics of McGraw and Wong. *Journal of Educational and Behavioral Statistics, 25*(2), 101–132.',
 'Yorke, M. (2006). *Employability in higher education: What it is – what it is not*. Higher Education Academy.',
 'Duoc UC, Escuela de Administración y Negocios, Sede San Bernardo (2026). *Proyecto La Puerta: Informe metodológico y análisis prospectivo de brechas en habilidades blandas para la práctica profesional* [documento interno].',
 'Duoc UC (2026). *Proyecto La Puerta: Base de datos y matriz de análisis (20261009_1658)* [base de datos interna, no publicada].',
]
def bibliografia(doc, subset=None):
    doc.part('REFERENCIAS')
    doc.p('Referencias en formato APA 7, ordenadas alfabéticamente. Las dos últimas son documentos internos del proyecto. Las referencias académicas fueron seleccionadas por la herramienta de IA a partir de su conocimiento de la literatura y **no se verificaron en línea durante esta sesión**: conviene confirmar volumen, páginas y DOI antes de una publicación formal.', 10, GRAY)
    for b in (BIB if subset is None else [BIB[i] for i in subset]):
        p = doc.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.15
        for part in re.split(r'(\*[^*]+\*)', b):
            if not part: continue
            it = part.startswith('*') and part.endswith('*'); r = p.add_run(part[1:-1] if it else part); r.italic = it; r.font.size = Pt(10.5); r.font.name = FONT

def declaracion_ia(doc, rol):
    doc.h2('Declaración sobre el uso de Inteligencia Artificial')
    doc.p('Este documento se elaboró con apoyo de **Claude (Anthropic)**, un modelo de Inteligencia Artificial generativa, entre octubre de 2026. A continuación se indica con transparencia qué hizo la herramienta y qué debe revisar una persona.', 10.5)
    doc.table(['Tarea', 'Qué hizo la IA', 'Qué debe revisar o decidir una persona'], [
     ['**Datos y cálculos**', 'Limpió la base y calculó promedios, porcentajes, comparaciones y gráficos con código (Python) guardado en la carpeta `src` del repositorio.', 'Revisar los scripts y repetir los cálculos ante cualquier duda. Los resultados son reproducibles.'],
     ['**Tarea específica de este documento**', rol[0].upper() + rol[1:] + '.', 'Validar que el enfoque y las conclusiones correspondan al propósito del proyecto.'],
     ['**Respuestas abiertas**', 'Las agrupó por temas de forma automática (clasificación preliminar).', 'Doble lectura humana y libro de códigos antes de citar conteos.'],
     ['**Redacción y diseño**', 'Redactó, organizó y diseñó tablas, gráficos, diagramas e infografías (ilustraciones propias, sin fotografías de personas).', 'Revisar el tono, el contenido y la adecuación institucional (por ejemplo, insertar el logo oficial).'],
     ['**Bibliografía**', 'Seleccionó las referencias a partir de su conocimiento de la literatura y, cuando se indica, de resúmenes encontrados en la web.', 'Verificar autores, volumen, páginas y DOI en la fuente antes de publicar. Lo no verificado está marcado.'],
     ['**Datos personales**', 'No se incluyeron nombres, correos ni códigos de estudiantes; las citas textuales son anónimas.', 'Mantener la confidencialidad al difundir.']], [3.4, 7.0, 6.1], size=8.8)
    doc.callout('Responsabilidad', ['La IA **no reemplaza el juicio experto** ni asume responsabilidad por las decisiones. Las interpretaciones y recomendaciones deben ser revisadas y asumidas por el equipo responsable del Proyecto La Puerta antes de su difusión institucional.', 'Si algún dato de este documento no coincide con las fuentes, prevalecen las fuentes originales.'])

def aviso_datos(doc):
    c = R['campos_vacios']
    doc.callout('Una advertencia importante sobre los datos',
        [f"La base trae **{R['N']} respuestas de un solo momento**. Los campos que permitirían comparar «antes y después» o «grupo La Puerta vs. grupo de comparación» (código de estudiante, momento de medición y grupo de asignación) vienen **vacíos en las {c['ID02 código']} filas**, igual que la pregunta de valoración del programa (POST_EXP1).",
         'Por eso este análisis describe **cómo se perciben hoy los estudiantes** y contrasta hipótesis sobre **diferencias entre perfiles** (por ejemplo, con y sin práctica previa). **No puede afirmar que el programa La Puerta produzca cambios**: eso requiere medir a las mismas personas antes y después, con un grupo de comparación.'])
